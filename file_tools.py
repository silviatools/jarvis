"""
file_tools.py — чтение и создание файлов популярных форматов для ассистента.

Чтение (read_file): txt/md/csv/json/html/код, docx, xlsx/xlsm, xls, pptx, pdf (с
текстовым слоем), rtf, odt/ods/odp, zip.
Создание (build): docx, xlsx, pptx, pdf (с кириллицей), а также csv/txt/md/json/html
как есть. Документы собираются из лёгкой markdown-разметки (заголовки #, списки,
таблицы |…|, **жирный**, *курсив*, `код`, ``` блоки ```, --- разделитель).

Тяжёлые библиотеки импортируются лениво: без них сервер работает, а нужный
формат отвечает понятной ошибкой.
"""

import io
import re
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path

FONT_DIR = Path(__file__).parent / "assets" / "fonts"

READ_LIMIT = 60000          # символов, которые отдаём модели за один файл
TEXT_EXT = {"txt", "md", "csv", "tsv", "json", "xml", "html", "htm", "yaml", "yml", "log", "ini", "py", "js",
            "ts", "sql", "sh", "ics", "srt", "css", "java", "c", "cpp", "h", "cs", "toml", "tex"}
READABLE_EXT = TEXT_EXT | {"docx", "xlsx", "xlsm", "xls", "pptx", "pdf", "rtf", "odt", "ods", "odp", "zip"}
BUILDABLE_EXT = {"docx", "xlsx", "pptx", "pdf"}   # из markdown/таблиц; остальное пишется как текст


class FileToolError(Exception):
    pass


def ext_of(name: str) -> str:
    return name.rsplit(".", 1)[-1].lower() if "." in (name or "") else ""


def _need(module: str, pkg: str):
    raise FileToolError(f"на сервере не установлена библиотека {pkg} (добавьте в requirements.txt и задеплойте)")


# ══ чтение ═══════════════════════════════════════════════════════════════════

def _decode(data: bytes) -> str:
    if data[:2] in (b"\xff\xfe", b"\xfe\xff"):
        return data.decode("utf-16", errors="replace")
    for enc in ("utf-8-sig", "cp1251"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _read_docx(data: bytes) -> str:
    try:
        import docx
    except ImportError:
        _need("docx", "python-docx")
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    d = docx.Document(io.BytesIO(data))
    out = []
    for child in d.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            p = Paragraph(child, d)
            t = p.text.strip()
            if not t:
                continue
            style = (p.style.name if p.style is not None else "") or ""
            m = re.match(r"Heading (\d)", style)
            if m:
                t = "#" * int(m.group(1)) + " " + t
            elif style == "Title":
                t = "# " + t
            elif "List" in style:
                t = "• " + t
            out.append(t)
        elif tag == "tbl":
            tbl = Table(child, d)
            for row in tbl.rows:
                seen, cells = set(), []
                for c in row.cells:       # объединённые ячейки повторяются — берём по разу
                    if id(c._tc) in seen:
                        continue
                    seen.add(id(c._tc))
                    cells.append(c.text.strip().replace("\n", " "))
                out.append("| " + " | ".join(cells) + " |")
    return "\n".join(out)


def _fmt_cell(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).replace("\n", " ")


def _read_xlsx(data: bytes) -> str:
    try:
        import openpyxl
    except ImportError:
        _need("openpyxl", "openpyxl")
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    wf = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=False)   # формулы без кэша значений
    out = []
    for ws in wb.worksheets:
        out.append(f"## Лист «{ws.title}»")
        n = 0
        frows = wf[ws.title].iter_rows(values_only=True)
        for row, frow in zip(ws.iter_rows(values_only=True), frows):
            cells = [_fmt_cell(v if v is not None or not (isinstance(f, str) and f.startswith("=")) else f)
                     for v, f in zip(row, frow)]
            while cells and not cells[-1]:
                cells.pop()
            if not cells:
                continue
            out.append(" | ".join(cells))
            n += 1
            if n >= 500:
                out.append("[…строки дальше обрезаны]")
                break
    return "\n".join(out)


def _read_xls(data: bytes) -> str:
    try:
        import xlrd
    except ImportError:
        _need("xlrd", "xlrd")
    wb = xlrd.open_workbook(file_contents=data)
    out = []
    for sh in wb.sheets():
        out.append(f"## Лист «{sh.name}»")
        for r in range(min(sh.nrows, 500)):
            cells = [_fmt_cell(sh.cell_value(r, c)) for c in range(sh.ncols)]
            while cells and not cells[-1]:
                cells.pop()
            if cells:
                out.append(" | ".join(cells))
    return "\n".join(out)


def _read_pptx(data: bytes) -> str:
    try:
        from pptx import Presentation
    except ImportError:
        _need("pptx", "python-pptx")
    prs = Presentation(io.BytesIO(data))
    out = []
    for n, slide in enumerate(prs.slides, 1):
        out.append(f"## Слайд {n}")
        for sh in slide.shapes:
            if sh.has_text_frame:
                for p in sh.text_frame.paragraphs:
                    t = "".join(r.text for r in p.runs).strip()
                    if t:
                        out.append(("  " * p.level) + "• " + t)
            elif getattr(sh, "has_table", False) and sh.has_table:
                for row in sh.table.rows:
                    out.append("| " + " | ".join(c.text.strip().replace("\n", " ") for c in row.cells) + " |")
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            note = slide.notes_slide.notes_text_frame.text.strip()
            if note:
                out.append(f"[Заметки докладчика: {note}]")
    return "\n".join(out)


def _read_pdf(data: bytes) -> str:
    try:
        from pypdf import PdfReader
    except ImportError:
        _need("pypdf", "pypdf")
    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted:
        try:
            if not reader.decrypt(""):
                raise FileToolError("PDF защищён паролем")
        except FileToolError:
            raise
        except Exception:
            raise FileToolError("PDF защищён паролем")
    out, total = [], 0
    for n, page in enumerate(reader.pages, 1):
        try:
            t = (page.extract_text() or "").strip()
        except Exception:
            t = ""
        if t:
            out.append(f"--- стр. {n} ---\n{t}")
            total += len(t)
        if total > READ_LIMIT * 2:
            out.append("[…остальные страницы пропущены]")
            break
    if not out:
        raise FileToolError("в PDF нет текстового слоя (похоже, скан) — распознавать картинки в PDF я не умею")
    return "\n".join(out)


def _read_rtf(data: bytes) -> str:
    s = data.decode("latin-1")
    s = re.sub(r"\{\\\*[^{}]*\}", "", s)           # служебные группы
    for _ in range(4):                              # вложенные группы шрифтов/стилей
        s = re.sub(r"\{\\(?:fonttbl|colortbl|stylesheet|info|pict)[^{}]*(?:\{[^{}]*\}[^{}]*)*\}", "", s)

    def hexch(m):
        return bytes([int(m.group(1), 16)]).decode("cp1251", errors="replace")

    def uni(m):
        n = int(m.group(1))
        return chr(n + 65536 if n < 0 else n)

    s = re.sub(r"\\'([0-9a-fA-F]{2})", hexch, s)
    s = re.sub(r"\\u(-?\d+)\??", uni, s)
    s = re.sub(r"\\(?:par|line)\b ?", "\n", s)
    s = re.sub(r"\\tab\b ?", "\t", s)
    s = re.sub(r"\\[a-zA-Z]+-?\d* ?", "", s)
    s = s.replace("\\{", "{").replace("\\}", "}").replace("\\\\", "\\")
    s = re.sub(r"[{}]", "", s)
    return re.sub(r"\n{3,}", "\n\n", s).strip()


def _odf_text(root: ET.Element) -> str:
    def local(t):
        return t.rsplit("}", 1)[-1]

    out = []

    def walk(el):
        name = local(el.tag)
        if name == "table-row":
            cells = []
            for c in el:
                if local(c.tag) in ("table-cell", "covered-table-cell"):
                    cells.append(" ".join("".join(p.itertext()) for p in c.iter() if local(p.tag) == "p").strip())
            while cells and not cells[-1]:
                cells.pop()
            if cells:
                out.append(" | ".join(cells))
            return
        if name in ("p", "h"):
            t = "".join(el.itertext()).strip()
            if t:
                out.append(("# " if name == "h" else "") + t)
            return
        if name == "table":
            out.append("## Лист/таблица «" + (el.attrib.get("{urn:oasis:names:tc:opendocument:xmlns:table:1.0}name") or "") + "»")
        for ch in el:
            walk(ch)

    walk(root)
    return "\n".join(out)


def _read_odf(data: bytes) -> str:
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        return _odf_text(ET.fromstring(z.read("content.xml")))


def _read_zip(data: bytes) -> str:
    out = []
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        infos = [i for i in z.infolist() if not i.is_dir()]
        out.append(f"Архив: {len(infos)} файл(ов)")
        for i in infos[:200]:
            out.append(f"• {i.filename} ({i.file_size} Б)")
        if len(infos) > 200:
            out.append("[…список обрезан]")
        budget = 40000
        for i in infos:
            if ext_of(i.filename) in TEXT_EXT and i.file_size <= 100000 and budget > 0:
                t = _decode(z.read(i))[:budget]
                budget -= len(t)
                out.append(f"\n=== {i.filename} ===\n{t}")
    return "\n".join(out)


_READERS = {"docx": _read_docx, "xlsx": _read_xlsx, "xlsm": _read_xlsx, "xls": _read_xls, "pptx": _read_pptx,
            "pdf": _read_pdf, "rtf": _read_rtf, "odt": _read_odf, "ods": _read_odf, "odp": _read_odf,
            "zip": _read_zip}
_LEGACY = {"doc": "docx", "ppt": "pptx"}


def read_file(data: bytes, name: str, limit: int = READ_LIMIT) -> dict:
    """→ {"text", "truncated", "format"}. FileToolError — понятная причина отказа."""
    ext = ext_of(name)
    if ext in _LEGACY:
        raise FileToolError(f"формат .{ext} устарел и не читается — пересохраните файл как .{_LEGACY[ext]}")
    try:
        if ext in _READERS:
            text = _READERS[ext](data)
        elif ext in TEXT_EXT or not ext:
            text = _decode(data)
        else:
            raise FileToolError(f"формат .{ext} не читаю")
    except FileToolError:
        raise
    except zipfile.BadZipFile:
        raise FileToolError("файл повреждён или это не тот формат")
    except Exception as e:
        raise FileToolError(f"не удалось прочитать файл: {e}")
    text = text.strip()
    truncated = len(text) > limit
    return {"text": text[:limit] + ("\n[…обрезано]" if truncated else ""), "truncated": truncated, "format": ext}


# ══ разбор markdown-разметки в блоки ═════════════════════════════════════════

_INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\*[^*\s][^*]*?\*|~~.+?~~)")


def parse_inline(text: str) -> list:
    """→ [(текст, {"b","i","code","s"})]"""
    runs = []
    for part in _INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            runs.append((part[2:-2], {"b"}))
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            runs.append((part[1:-1], {"code"}))
        elif part.startswith("~~") and part.endswith("~~") and len(part) > 4:
            runs.append((part[2:-2], {"s"}))
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            runs.append((part[1:-1], {"i"}))
        else:
            runs.append((part, set()))
    return runs


def plain(text: str) -> str:
    return "".join(t for t, _ in parse_inline(text))


def _table_cells(line: str) -> list:
    s = line.strip()
    if s.startswith("|"):
        s = s[1:]
    if s.endswith("|"):
        s = s[:-1]
    return [c.strip() for c in s.split("|")]


def parse_blocks(src: str) -> list:
    """Блоки: ("heading", (lvl, text)), ("para", text), ("bullet", (lvl, text)), ("number", (n, text)),
    ("table", rows), ("code", text), ("quote", text), ("hr", None)."""
    lines = (src or "").replace("\r\n", "\n").replace("\r", "\n").split("\n")
    blocks, i = [], 0
    special = re.compile(r"^\s*(#{1,6}\s|[-*•]\s|\d+[.)]\s|>|\||```|-{3,}\s*$|\*{3,}\s*$|_{3,}\s*$)")
    while i < len(lines):
        ln = lines[i]
        s = ln.strip()
        if not s:
            i += 1
            continue
        if s.startswith("```"):
            i += 1
            buf = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                buf.append(lines[i])
                i += 1
            i += 1
            blocks.append(("code", "\n".join(buf)))
            continue
        if re.fullmatch(r"(-{3,}|\*{3,}|_{3,})", s):
            blocks.append(("hr", None))
            i += 1
            continue
        m = re.match(r"(#{1,6})\s+(.*)", s)
        if m:
            blocks.append(("heading", (len(m.group(1)), m.group(2).strip())))
            i += 1
            continue
        if s.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = _table_cells(lines[i])
                if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c) or not any(cells):
                    rows.append(cells)
                i += 1
            if rows:
                w = max(len(r) for r in rows)
                blocks.append(("table", [r + [""] * (w - len(r)) for r in rows]))
            continue
        m = re.match(r"^(\s*)[-*•]\s+(.*)", ln)
        if m:
            t = m.group(2)
            t = re.sub(r"^\[ \]\s*", "☐ ", t)
            t = re.sub(r"^\[[xX]\]\s*", "☑ ", t)
            blocks.append(("bullet", (min(len(m.group(1).replace("\t", "  ")) // 2, 3), t)))
            i += 1
            continue
        m = re.match(r"^\s*(\d+)[.)]\s+(.*)", ln)
        if m:
            blocks.append(("number", (int(m.group(1)), m.group(2))))
            i += 1
            continue
        if s.startswith(">"):
            buf = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                buf.append(lines[i].strip().lstrip(">").strip())
                i += 1
            blocks.append(("quote", "\n".join(buf)))
            continue
        buf = []
        while i < len(lines) and lines[i].strip() and not (buf and special.match(lines[i])):
            buf.append(lines[i].strip())
            i += 1
        blocks.append(("para", "\n".join(buf)))
    return blocks


# ══ DOCX ═════════════════════════════════════════════════════════════════════

def _docx_runs(par, text: str, base_italic=False, size=None):
    from docx.shared import Pt
    lines = text.split("\n")
    for li, line in enumerate(lines):
        for t, fl in parse_inline(line):
            r = par.add_run(t)
            r.bold = "b" in fl or None
            r.italic = ("i" in fl or base_italic) or None
            r.font.strike = "s" in fl or None
            if "code" in fl:
                r.font.name = "Courier New"
            if size:
                r.font.size = Pt(size)
        if li < len(lines) - 1:
            par.add_run().add_break()


def _shade(cell, color="D9E2F3"):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    tcPr.append(shd)


def _hr_docx(d):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    p = d.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    b = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for k, v in (("val", "single"), ("sz", "6"), ("space", "1"), ("color", "999999")):
        bottom.set(qn(f"w:{k}"), v)
    b.append(bottom)
    pPr.append(b)


def build_docx(blocks: list) -> bytes:
    try:
        import docx
        from docx.shared import Cm, Pt
    except ImportError:
        _need("docx", "python-docx")
    d = docx.Document()
    st = d.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(11)
    for kind, val in blocks:
        if kind == "heading":
            lvl, text = val
            d.add_heading(plain(text), level=min(lvl, 4))
        elif kind == "para":
            _docx_runs(d.add_paragraph(), val)
        elif kind == "bullet":
            lvl, text = val
            style = "List Bullet" if lvl == 0 else f"List Bullet {min(lvl + 1, 3)}"
            _docx_runs(d.add_paragraph(style=style), text)
        elif kind == "number":
            n, text = val
            p = d.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.9)
            p.paragraph_format.first_line_indent = Cm(-0.6)
            p.add_run(f"{n}. ")
            _docx_runs(p, text)
        elif kind == "quote":
            p = d.add_paragraph()
            p.paragraph_format.left_indent = Cm(1)
            _docx_runs(p, val, base_italic=True)
        elif kind == "code":
            p = d.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.5)
            r = p.add_run(val)
            r.font.name = "Courier New"
            r.font.size = Pt(9.5)
        elif kind == "table":
            t = d.add_table(rows=len(val), cols=len(val[0]))
            t.style = "Table Grid"
            for ri, row in enumerate(val):
                for ci, cell_text in enumerate(row):
                    cell = t.cell(ri, ci)
                    cell.text = ""
                    _docx_runs(cell.paragraphs[0], cell_text)
                    if ri == 0:
                        for r in cell.paragraphs[0].runs:
                            r.bold = True
                        _shade(cell)
            d.add_paragraph()
        elif kind == "hr":
            _hr_docx(d)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


# ══ XLSX ═════════════════════════════════════════════════════════════════════

_NUM = re.compile(r"^-?\d+(?:[.,]\d+)?$")


def _cell_value(v):
    """Числа-строки → числа (кроме «007», телефонов и кодов с нулём в начале)."""
    if isinstance(v, str):
        s = v.strip()
        if _NUM.match(s) and not (len(s.lstrip("-")) > 1 and s.lstrip("-").startswith("0") and not s.lstrip("-").startswith(("0.", "0,"))) \
                and len(s) < 16:
            f = float(s.replace(",", "."))
            return int(f) if f.is_integer() and "." not in s and "," not in s else f
        return v
    return v


def build_xlsx(sheets: list) -> bytes:
    try:
        import openpyxl
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        _need("openpyxl", "openpyxl")
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    used = set()
    for idx, sh in enumerate(sheets, 1):
        name = re.sub(r"[\\/*?:\[\]]", " ", str(sh.get("name") or f"Лист{idx}")).strip()[:31] or f"Лист{idx}"
        base, k = name, 2
        while name.lower() in used:
            name = f"{base[:28]}_{k}"
            k += 1
        used.add(name.lower())
        ws = wb.create_sheet(name)
        rows = sh.get("rows") or []
        for r in rows:
            ws.append([_cell_value(v) for v in r])
        if rows:
            for c in ws[1]:
                c.font = Font(bold=True)
                c.fill = PatternFill("solid", fgColor="D9E2F3")
                c.alignment = Alignment(vertical="center", wrap_text=True)
            ws.freeze_panes = "A2"
            ncols = max(len(r) for r in rows)
            for ci in range(1, ncols + 1):
                width = max(len(str(r[ci - 1])) if ci <= len(r) and r[ci - 1] is not None else 0 for r in rows)
                ws.column_dimensions[get_column_letter(ci)].width = min(max(10, width + 2), 60)
    if not wb.worksheets:
        wb.create_sheet("Лист1")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def sheets_from_content(content: str) -> list:
    """Таблица из текста: markdown-таблицы (несколько → несколько листов) или CSV/TSV."""
    blocks = parse_blocks(content)
    tables = [(v) for k, v in blocks if k == "table"]
    if tables:
        names = []
        last_head = None
        for k, v in blocks:
            if k == "heading":
                last_head = plain(v[1])
            elif k == "table":
                names.append(last_head)
                last_head = None
        return [{"name": names[i] or f"Лист{i + 1}", "rows": t} for i, t in enumerate(tables)]
    import csv
    sample = content[:2000]
    delim = ";" if sample.count(";") > sample.count(",") and sample.count(";") >= sample.count("\t") else \
        ("\t" if sample.count("\t") > sample.count(",") else ",")
    rows = list(csv.reader(io.StringIO(content.strip()), delimiter=delim))
    return [{"name": "Лист1", "rows": rows}]


# ══ PPTX ═════════════════════════════════════════════════════════════════════

def build_pptx(blocks: list) -> bytes:
    try:
        from pptx import Presentation
        from pptx.util import Inches, Pt
    except ImportError:
        _need("pptx", "python-pptx")
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    W = 13.333

    # Слайд начинается на заголовках уровня 1–2 и на разделителе «---».
    slides, cur = [], None
    for kind, val in blocks:
        if kind == "hr":
            cur = None
            continue
        if kind == "heading" and val[0] <= 2:
            cur = {"title": plain(val[1]), "items": []}
            slides.append(cur)
            continue
        if cur is None:
            cur = {"title": "", "items": []}
            slides.append(cur)
        cur["items"].append((kind, val))
    if not slides:
        slides = [{"title": "Презентация", "items": []}]

    for si, sl in enumerate(slides):
        tables = [v for k, v in sl["items"] if k == "table"]
        only_text = not tables and all(k == "para" for k, _ in sl["items"])
        title_slide = si == 0 and sl["title"] and only_text and len(sl["items"]) <= 2
        slide = prs.slides.add_slide(prs.slide_layouts[0 if title_slide else 1])
        title_ph = slide.shapes.title
        title_ph.text = sl["title"] or " "
        title_ph.left, title_ph.top, title_ph.width, title_ph.height = (
            Inches(0.7), Inches(2.3 if title_slide else 0.4), Inches(W - 1.4), Inches(1.5 if title_slide else 1.2))
        for p in title_ph.text_frame.paragraphs:
            for r in p.runs:
                r.font.size = Pt(44 if title_slide else 34)
                r.font.bold = True
        body = slide.placeholders[1]
        if title_slide:
            body.left, body.top, body.width, body.height = Inches(0.7), Inches(4.0), Inches(W - 1.4), Inches(2.0)
            body.text_frame.text = "\n".join(plain(v) for _, v in sl["items"])
            continue
        if tables:
            body._element.getparent().remove(body._element)
            rows = tables[0]
            gt = slide.shapes.add_table(len(rows), len(rows[0]), Inches(0.7), Inches(1.8), Inches(W - 1.4),
                                        Inches(min(0.5 * len(rows), 5.2))).table
            for ri, row in enumerate(rows):
                for ci, txt in enumerate(row):
                    cell = gt.cell(ri, ci)
                    cell.text = plain(txt)
                    for p in cell.text_frame.paragraphs:
                        for r in p.runs:
                            r.font.size = Pt(16)
                            r.font.bold = (ri == 0) or None
            continue
        body.left, body.top, body.width, body.height = Inches(0.7), Inches(1.8), Inches(W - 1.4), Inches(5.3)
        tf = body.text_frame
        tf.word_wrap = True
        n_lines = sum(max(1, len(str(v if isinstance(v, str) else v[1])) // 90 + 1) for _, v in sl["items"])
        size = 28 if n_lines <= 5 else 24 if n_lines <= 8 else 20 if n_lines <= 12 else 16 if n_lines <= 17 else 13
        first = True
        for kind, val in sl["items"]:
            par = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            level, text, bold = 0, "", False
            if kind == "bullet":
                level, text = val
            elif kind == "number":
                text = f"{val[0]}. {val[1]}"
            elif kind == "heading":
                text, bold = plain(val[1]), True
            elif kind == "code":
                text = val
            else:
                text = val if isinstance(val, str) else str(val)
            par.level = min(level, 4)
            for t, fl in parse_inline(text):
                r = par.add_run()
                r.text = t
                r.font.size = Pt(size - 2 * min(level, 2))
                r.font.bold = bold or ("b" in fl) or None
                r.font.italic = ("i" in fl) or None
                if kind == "code" or "code" in fl:
                    r.font.name = "Courier New"
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


# ══ PDF ══════════════════════════════════════════════════════════════════════

def build_pdf(blocks: list) -> bytes:
    try:
        from fpdf import FPDF
        from fpdf.fonts import FontFace
    except ImportError:
        _need("fpdf", "fpdf2")
    sans, bold, mono = FONT_DIR / "DejaVuSans.ttf", FONT_DIR / "DejaVuSans-Bold.ttf", FONT_DIR / "DejaVuSansMono.ttf"
    if not sans.exists():
        raise FileToolError("в проекте нет шрифтов assets/fonts (DejaVu) — PDF с кириллицей не собрать")
    pdf = FPDF(format="A4")
    pdf.set_margins(20, 18, 20)
    pdf.set_auto_page_break(True, 18)
    pdf.add_font("D", "", str(sans))
    pdf.add_font("D", "B", str(bold))
    pdf.add_font("M", "", str(mono))
    pdf.add_page()

    def runs(text, size=11, lh=6.0, indent=0.0, prefix=""):
        old = pdf.l_margin
        pdf.l_margin = old + indent
        pdf.set_x(pdf.l_margin)
        pdf.set_font("D", "", size)
        if prefix:
            pdf.write(lh, prefix)
        for line in text.split("\n"):
            for t, fl in parse_inline(line):
                if "code" in fl:
                    pdf.set_font("M", "", size - 1)
                else:
                    pdf.set_font("D", "B" if "b" in fl else "", size)
                pdf.write(lh, t)
            pdf.ln(lh)
        pdf.l_margin = old
        pdf.set_x(old)

    sizes = {1: 20, 2: 16, 3: 13.5, 4: 12}
    for kind, val in blocks:
        if kind == "heading":
            lvl, text = val
            pdf.ln(2)
            pdf.set_font("D", "B", sizes.get(lvl, 12))
            pdf.multi_cell(0, sizes.get(lvl, 12) * 0.55 + 2, plain(text), align="L", new_x="LMARGIN", new_y="NEXT")
            pdf.ln(1)
        elif kind == "para":
            runs(val)
            pdf.ln(1.5)
        elif kind == "bullet":
            lvl, text = val
            runs(text, indent=5 + lvl * 6, prefix="•  ")
        elif kind == "number":
            n, text = val
            runs(text, indent=5, prefix=f"{n}.  ")
        elif kind == "quote":
            pdf.set_text_color(90, 90, 90)
            runs(val, indent=8)
            pdf.set_text_color(0, 0, 0)
            pdf.ln(1.5)
        elif kind == "code":
            pdf.set_font("M", "", 9)
            pdf.set_fill_color(240, 240, 240)
            pdf.multi_cell(0, 4.6, val, align="L", fill=True, new_x="LMARGIN", new_y="NEXT")
            pdf.ln(2)
        elif kind == "hr":
            y = pdf.get_y() + 2
            pdf.set_draw_color(170, 170, 170)
            pdf.line(pdf.l_margin, y, pdf.w - pdf.r_margin, y)
            pdf.ln(5)
        elif kind == "table":
            pdf.set_font("D", "", 10)
            with pdf.table(text_align="LEFT", line_height=5.5, first_row_as_headings=True,
                           headings_style=FontFace(emphasis="BOLD", fill_color=(217, 226, 243))) as tb:
                for row in val:
                    r = tb.row()
                    for c in row:
                        r.cell(plain(c))
            pdf.ln(3)
    return bytes(pdf.output())


# ══ единая точка входа ═══════════════════════════════════════════════════════

def build(filename: str, content: str = None, sheets: list = None) -> tuple:
    """→ (bytes, filename). Формат определяется расширением filename."""
    filename = re.sub(r"[\\/\x00]", "_", (filename or "").strip())
    ext = ext_of(filename)
    if not ext:
        raise FileToolError("в имени файла нужно расширение (например report.docx)")
    if sheets is not None and not isinstance(sheets, list):
        raise FileToolError("sheets должен быть списком [{name, rows}]")
    if ext == "xlsx":
        if sheets:
            return build_xlsx(sheets), filename
        if content:
            return build_xlsx(sheets_from_content(content)), filename
        raise FileToolError("для xlsx нужен sheets [{name, rows}] или content (таблица/CSV)")
    if content is None or not str(content).strip():
        raise FileToolError("пустое содержимое файла")
    content = str(content)
    if ext in ("docx", "pdf", "pptx"):
        blocks = parse_blocks(content)
        return {"docx": build_docx, "pdf": build_pdf, "pptx": build_pptx}[ext](blocks), filename
    data = content.encode("utf-8")
    if ext in ("csv", "tsv"):
        data = b"\xef\xbb\xbf" + data       # BOM — чтобы Excel не ломал кириллицу
    return data, filename
